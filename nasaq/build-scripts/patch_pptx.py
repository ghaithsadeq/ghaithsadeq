# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Emu

NAME = "غيث صادق كيوان"
OLD, NEW = "GEN CODE TEAM", "إعداد وتقديم: " + NAME + " · GEN CODE TEAM"
BRAND_OLD, BRAND_NEW = "GEN CODE TEAM", NAME + " · GEN CODE TEAM"

def sub_runs(tf, old, new):
    n = 0
    for pa in tf.paragraphs:
        for r in pa.runs:
            if old in r.text:
                r.text = r.text.replace(old, new); n += 1
    return n

for path in ("build/final/Nasaq-v2-Presentation.pptx", "build/phase1/Nasaq-Presentation.pptx"):
    pr = Presentation(path)
    slides = layouts = 0
    for s in pr.slides:                       # cover + closing credit lines
        for sh in s.shapes:
            if sh.has_text_frame and OLD in sh.text_frame.text:
                slides += sub_runs(sh.text_frame, OLD, NEW)
    for l in pr.slide_layouts:                # running brand on every inner slide
        for sh in l.shapes:
            if sh.has_text_frame and BRAND_OLD in sh.text_frame.text:
                if sub_runs(sh.text_frame, BRAND_OLD, BRAND_NEW):
                    sh.width = Emu(4389120)   # widen so the longer line never wraps
                    sh.text_frame.word_wrap = False
                    layouts += 1
    pr.save(path)
    assert slides and layouts, path
    print("%-34s %d slide credits · %d layout brands · %d slides"
          % (os.path.basename(path), slides, layouts, len(pr.slides._sldIdLst)))
